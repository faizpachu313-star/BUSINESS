#!/usr/bin/env python3
"""Build the AED <-> INR daily exchange tracker workbook.

One sheet per day (e.g. "01-Jan") plus a "Monthly Summary" sheet. Every
number on the sheets is a live Excel formula; the only things this script
writes as plain values are the inputs: the daily rate and the transactions
(Name, AED Received, INR Given) read from data/transactions.json.

Usage:
    python build_workbook.py [--data data/transactions.json] [--out OUTFILE]
    python scripts/recalc.py OUTFILE        # fills cached values (xlsx skill)
"""
import argparse
import calendar
import json
import os
import shutil
import subprocess
import tempfile
from datetime import date

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink

from compact_xlsx import compact

FONT = "Arial"
MIN_ROWS = 40          # pre-built formula rows per daily sheet (grows if a day needs more)
EXTRA_ROWS = 10        # spare rows kept below the last transaction of a busy day

# ---- number formats --------------------------------------------------------
AED_FMT = '#,##0.00" د.إ"'                      # amount first, Arabic symbol after (reads LTR everywhere)
INR_FMT = '"₹"#,##0.00'
PL_FMT = '"+₹"#,##0.00;"−₹"#,##0.00;"₹"0.00'            # +profit / −loss / zero
PROFIT_FMT = '"+₹"#,##0.00;"−₹"#,##0.00;"–"'
SUM_AED_FMT = '#,##0.00" د.إ";-#,##0.00" د.إ";"–"'
SUM_INR_FMT = '"₹"#,##0.00;"-₹"#,##0.00;"–"'
COUNT_FMT = '0;-0;"–"'
RATE_FMT = '"₹ "0.00##'

# ---- palette ---------------------------------------------------------------
NAVY, BLUE, GREY_FILL, INPUT_FILL = "1F3A5F", "2F5597", "F2F2F2", "FFF2CC"
LINE = "D0D0D0"
GREEN_TXT, GREEN_FILL, RED_TXT, RED_FILL, NEUTRAL_TXT = "006100", "C6EFCE", "9C0006", "FFC7CE", "7F7F7F"

thin = Side(style="thin", color=LINE)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def font(size=11, bold=False, color="000000", italic=False, underline=None):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic, underline=underline)


def fill(hex_color):
    return PatternFill("solid", start_color=hex_color, end_color=hex_color)


def sheet_name(day, mon_abbr):
    return f"{day:02d}-{mon_abbr}"


def add_sign_colouring(ws, rng, first_cell):
    """Green for > 0, red for < 0, grey for 0 (blank text results are ignored)."""
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f"AND(ISNUMBER({first_cell}),{first_cell}>0)"],
        font=Font(name=FONT, color=GREEN_TXT, bold=True), fill=fill(GREEN_FILL)))
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f"AND(ISNUMBER({first_cell}),{first_cell}<0)"],
        font=Font(name=FONT, color=RED_TXT, bold=True), fill=fill(RED_FILL)))
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f"AND(ISNUMBER({first_cell}),{first_cell}=0)"],
        font=Font(name=FONT, color=NEUTRAL_TXT)))


def build_day_sheet(wb, name, day_data):
    """Returns the row numbers of the totals block: dict label -> row."""
    txns = day_data.get("transactions", [])
    n_rows = max(MIN_ROWS, len(txns) + EXTRA_ROWS)
    HEADER, FIRST = 4, 5
    LAST = FIRST + n_rows - 1
    rng = lambda col: f"{col}{FIRST}:{col}{LAST}"

    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDE", (17, 30, 21, 21, 21)):
        ws.column_dimensions[col].width = w

    # Row 1: currency banner
    ws.merge_cells("A1:C1")
    ws.merge_cells("D1:E1")
    ws["A1"] = "🇦🇪 AED — UAE Dirham"
    ws["D1"] = "🇮🇳 INR — Indian Rupee"
    for c in ("A1", "B1", "C1", "D1", "E1"):
        ws[c].fill = fill(NAVY)
        ws[c].font = font(14, True, "FFFFFF")
        ws[c].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32

    # Row 2: exchange-rate input  ->  "1 AED = ₹ 25.00"
    ws["A2"] = "EXCHANGE RATE"
    ws["B2"] = "1 AED  ="
    rate_cell = ws["C2"]
    rate_cell.value = day_data.get("rate")
    rate_cell.number_format = RATE_FMT
    ws["D2"] = '=IF(N(C2)>0,"","◄ enter today\'s rate")'
    for c in ("A2", "B2", "C2", "D2", "E2"):
        ws[c].fill = fill(GREY_FILL)
    ws["A2"].font = font(10, True, NAVY)
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws["B2"].font = font(14, True, NAVY)
    ws["B2"].alignment = Alignment(horizontal="right", vertical="center")
    rate_cell.font = font(16, True, "0000FF")
    rate_cell.fill = fill(INPUT_FILL)
    rate_cell.alignment = Alignment(horizontal="center", vertical="center")
    med = Side(style="medium", color="BF9000")
    rate_cell.border = Border(left=med, right=med, top=med, bottom=med)
    ws["D2"].font = font(9, False, RED_TXT, italic=True)
    ws["D2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[2].height = 34
    dv_rate = DataValidation(type="decimal", operator="greaterThan", formula1="0", allow_blank=True,
                             showErrorMessage=True, errorTitle="Rate", error="Enter a number above 0, e.g. 25.00")
    ws.add_data_validation(dv_rate)
    dv_rate.add("C2")

    # Row 3: legend + worked example
    ws.merge_cells("A3:E3")
    ws["A3"] = ("Yellow cell and blue figures are yours to type; everything else calculates automatically. "
                "A blank Rate uses today's rate above.\n"
                "Profit / Loss = AED Received × Rate − INR Given   (100 AED at 25, ₹2,400 given → +₹100)")
    ws["A3"].font = font(9, False, NEUTRAL_TXT, italic=True)
    ws["A3"].alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
    ws.row_dimensions[3].height = 28

    # Row 4: table header (frozen)
    for col, label in zip("ABCDE", ("Rate", "Name", "AED Received", "INR Given", "Profit / Loss")):
        c = ws[f"{col}{HEADER}"]
        c.value = label
        c.font = font(11, True, "FFFFFF")
        c.fill = fill(BLUE)
        c.alignment = Alignment(horizontal="left" if col == "B" else "right",
                                vertical="center", indent=1)
        c.border = BOX
    ws.row_dimensions[HEADER].height = 24

    # Transaction rows: inputs in A:D, formula in E. A row's own rate wins over today's rate in C2.
    for i in range(n_rows):
        r = FIRST + i
        t = txns[i] if i < len(txns) else {}
        ws[f"A{r}"] = t.get("rate")
        ws[f"B{r}"] = t.get("name")
        ws[f"C{r}"] = t.get("aed")
        ws[f"D{r}"] = t.get("inr")
        rate = f"IF(ISNUMBER(A{r}),A{r},$C$2)"
        ws[f"E{r}"] = (f'=IF(AND(ISNUMBER(C{r}),ISNUMBER(D{r}),N({rate})>0),'
                       f'ROUND(C{r}*{rate}-D{r},2),"")')
        ws[f"B{r}"].font = font(11)
        ws[f"B{r}"].alignment = Alignment(horizontal="left", indent=1)
        for col, fmt in (("A", RATE_FMT), ("C", AED_FMT), ("D", INR_FMT)):
            ws[f"{col}{r}"].number_format = fmt
            ws[f"{col}{r}"].font = font(11, color="0000FF")
            ws[f"{col}{r}"].alignment = Alignment(horizontal="right", readingOrder=1)
        ws[f"E{r}"].number_format = PL_FMT
        ws[f"E{r}"].font = font(11, True)
        ws[f"E{r}"].alignment = Alignment(horizontal="right")
        for col in "ABCDE":
            ws[f"{col}{r}"].border = BOX

    dv_amt = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0",
                            allow_blank=True, showErrorMessage=True, errorTitle="Amount",
                            error="Numbers only, 0 or more")
    ws.add_data_validation(dv_amt)
    dv_amt.add(f"C{FIRST}:D{LAST}")
    dv_row_rate = DataValidation(type="decimal", operator="greaterThan", formula1="0", allow_blank=True,
                                 showErrorMessage=True, errorTitle="Rate",
                                 error="Enter a number above 0, or leave blank for today's rate")
    ws.add_data_validation(dv_row_rate)
    dv_row_rate.add(f"A{FIRST}:A{LAST}")

    # Daily totals
    T = LAST + 2
    ws.merge_cells(f"A{T}:E{T}")
    ws[f"A{T}"] = "DAILY TOTALS"
    ws[f"A{T}"].font = font(12, True, "FFFFFF")
    ws[f"A{T}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for col in "ABCDE":
        ws[f"{col}{T}"].fill = fill(NAVY)
    ws.row_dimensions[T].height = 24

    rows = {}
    items = [
        ("aed", "Total AED Received", f"=SUM({rng('C')})", AED_FMT),
        ("inr", "Total INR Given", f"=SUM({rng('D')})", INR_FMT),
        ("profit", "Total Profit", f'=SUMIF({rng("E")},">0")', PL_FMT),
        ("loss", "Total Loss", f'=SUMIF({rng("E")},"<0")', PL_FMT),
        ("net", "Net Profit/Loss", f"=SUM({rng('E')})", PL_FMT),
        ("count", "Number of Transactions",
         f"=SUMPRODUCT(--((ISNUMBER({rng('C')})+ISNUMBER({rng('D')}))>0))", "0"),
    ]
    for k, (key, label, formula, fmt) in enumerate(items, start=1):
        r = T + k
        rows[key] = r
        ws.merge_cells(f"A{r}:D{r}")
        ws[f"A{r}"] = label
        ws[f"A{r}"].font = font(11, True)
        ws[f"A{r}"].alignment = Alignment(horizontal="right", indent=1)
        ws[f"E{r}"] = formula
        ws[f"E{r}"].number_format = fmt
        ws[f"E{r}"].font = font(11, True)
        ws[f"E{r}"].alignment = Alignment(horizontal="right", readingOrder=1)
        for col in "ABCDE":
            ws[f"{col}{r}"].border = BOX
    ws[f"E{rows['profit']}"].font = font(11, True, GREEN_TXT)
    ws[f"E{rows['loss']}"].font = font(11, True, RED_TXT)
    for col in "ABCDE":
        ws[f"{col}{rows['net']}"].border = Border(left=thin, right=thin, bottom=thin,
                                                  top=Side(style="medium", color=NAVY))
    # the formula is written for E5 and shifts for every other cell in the range
    add_sign_colouring(ws, f"E{FIRST}:E{LAST} E{rows['net']}", f"$E{FIRST}")

    ws.freeze_panes = f"A{FIRST}"
    ws.print_title_rows = f"1:{HEADER}"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return rows


def build_summary(wb, names, totals_rows, month_name):
    ws = wb["Monthly Summary"]
    ws.sheet_properties.tabColor = NAVY
    ws.sheet_view.showGridLines = False
    widths = (14, 21, 21, 19, 19, 21, 16)
    for col, w in zip("ABCDEFG", widths):
        ws.column_dimensions[col].width = w

    ws.merge_cells("A1:G1")
    ws["A1"] = f"🇦🇪 AED  ⇄  🇮🇳 INR  —  Monthly Summary ({month_name})"
    ws["A1"].font = font(15, True, "FFFFFF")
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for col in "ABCDEFG":
        ws[f"{col}1"].fill = fill(NAVY)
    ws.row_dimensions[1].height = 34

    ws.merge_cells("A2:G2")
    ws["A2"] = ("Pulled automatically from the daily sheets · losses are shown as negative "
                "amounts · Net = Profit + Loss · click a day to open its sheet")
    ws["A2"].font = font(9, False, NEUTRAL_TXT, italic=True)
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    HEADER, FIRST = 4, 5
    heads = ("Day", "AED Received", "INR Given", "Profit", "Loss", "Net Profit/Loss", "Transactions")
    for col, label in zip("ABCDEFG", heads):
        c = ws[f"{col}{HEADER}"]
        c.value = label
        c.font = font(11, True, "FFFFFF")
        c.fill = fill(BLUE)
        c.alignment = Alignment(horizontal="left" if col == "A" else "right",
                                vertical="center", indent=1)
        c.border = BOX
    ws.row_dimensions[HEADER].height = 24

    src = (("B", "aed", SUM_AED_FMT), ("C", "inr", SUM_INR_FMT), ("D", "profit", PROFIT_FMT),
           ("E", "loss", PROFIT_FMT), ("F", "net", PL_FMT), ("G", "count", COUNT_FMT))
    for i, name in enumerate(names):
        r = FIRST + i
        ws[f"A{r}"] = name
        ws[f"A{r}"].hyperlink = Hyperlink(ref=f"A{r}", location=f"'{name}'!A1")
        ws[f"A{r}"].font = font(11, color="0563C1", underline="single")
        ws[f"A{r}"].alignment = Alignment(horizontal="left", indent=1)
        for col, key, fmt in src:
            c = ws[f"{col}{r}"]
            c.value = f"='{name}'!E{totals_rows[name][key]}"
            c.number_format = fmt
            c.font = font(11, bold=(col == "F"))
            c.alignment = Alignment(horizontal="right", readingOrder=1)
        for col in "ABCDEFG":
            ws[f"{col}{r}"].border = BOX
        ws[f"D{r}"].font = font(11, color=GREEN_TXT)
        ws[f"E{r}"].font = font(11, color=RED_TXT)
    LAST = FIRST + len(names) - 1

    # Month totals
    T = LAST + 1
    ws[f"A{T}"] = "MONTH TOTAL"
    for col, key, fmt in src:
        ws[f"{col}{T}"] = f"=SUM({col}{FIRST}:{col}{LAST})"
        ws[f"{col}{T}"].number_format = {"B": AED_FMT, "C": INR_FMT, "D": PL_FMT,
                                         "E": PL_FMT, "F": PL_FMT, "G": "0"}[col]
    top = Side(style="medium", color=NAVY)
    for col in "ABCDEFG":
        c = ws[f"{col}{T}"]
        c.font = font(12, True)
        c.fill = fill(GREY_FILL)
        c.border = Border(left=thin, right=thin, bottom=thin, top=top)
        c.alignment = Alignment(horizontal="left" if col == "A" else "right",
                                vertical="center", indent=1 if col == "A" else 0, readingOrder=1)
    ws[f"D{T}"].font = font(12, True, GREEN_TXT)
    ws[f"E{T}"].font = font(12, True, RED_TXT)
    ws.row_dimensions[T].height = 26

    add_sign_colouring(ws, f"F{FIRST}:F{LAST}", f"$F{FIRST}")
    add_sign_colouring(ws, f"F{T}", f"$F${T}")
    ws.freeze_panes = f"B{FIRST}"
    ws.print_title_rows = f"1:{HEADER}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def libreoffice_recalc(raw_path, workdir):
    """Return a LibreOffice-calculated copy of raw_path, or None if LibreOffice is unavailable."""
    if not shutil.which("soffice"):
        return None
    outdir = os.path.join(workdir, "lo")
    os.makedirs(outdir, exist_ok=True)
    env = dict(os.environ, SAL_USE_VCLPLUGIN="svp")
    subprocess.run(["soffice", f"-env:UserInstallation=file://{workdir}/profile", "--headless",
                    "--norestore", "--convert-to", "xlsx", "--outdir", outdir, raw_path],
                   env=env, check=True, capture_output=True, timeout=300)
    return os.path.join(outdir, os.path.basename(raw_path))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--data", default=os.path.join(os.path.dirname(__file__), "data", "transactions.json"))
    ap.add_argument("--out", default=None)
    args = ap.parse_args()

    data = {}
    if os.path.exists(args.data):
        with open(args.data, encoding="utf-8") as fh:
            data = json.load(fh)
    month = int(data.get("month") or 1)
    year = int(data.get("year") or date.today().year)
    mon_abbr = calendar.month_abbr[month]
    n_days = calendar.monthrange(year, month)[1]
    days = data.get("days", {})

    # a data file that names its year gets a per-year file name, so months don't collide across years
    label = f"{mon_abbr} {year}" if data.get("year") else mon_abbr
    out = args.out or os.path.join(os.path.dirname(__file__), "output",
                                   f"AED-INR Exchange Tracker - {label}.xlsx")
    os.makedirs(os.path.dirname(out), exist_ok=True)

    wb = Workbook()
    wb.active.title = "Monthly Summary"          # first tab; filled in after the day sheets exist
    names, totals_rows = [], {}
    for d in range(1, n_days + 1):
        name = sheet_name(d, mon_abbr)
        totals_rows[name] = build_day_sheet(wb, name, days.get(str(d), days.get(f"{d:02d}", {})))
        names.append(name)
    build_summary(wb, names, totals_rows, calendar.month_name[month])
    wb.active = 0
    with tempfile.TemporaryDirectory() as tmp:
        raw = os.path.join(tmp, "raw.xlsx")
        wb.save(raw)
        calculated = libreoffice_recalc(raw, tmp)
        if calculated is None:
            print("warning: LibreOffice not found - saving without cached formula results")
        compact(raw, out, calculated)
    print(out, f"{os.path.getsize(out) / 1024:.1f} KB")


if __name__ == "__main__":
    main()
