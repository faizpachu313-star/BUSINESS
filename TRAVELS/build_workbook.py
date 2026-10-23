#!/usr/bin/env python3
"""Build the Faiz Travels monthly workbook in the Faiz Medical Records house style.

A "Summary" sheet followed by one sheet per month (``Jan-2026`` ...). Each month
sheet has Sl No / Date / Name / Service / Charged, blank rows ready to fill,
and totals at the bottom (number of entries and the month's total Charged).
Sl No, totals and the Summary are live formulas.

Usage:
    python build_workbook.py                 # 2026
    python build_workbook.py --year 2027 [--rows 60] [--out FILE.xlsx]
"""
import argparse
import calendar

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

# ---- Faiz Medical Records tokens ------------------------------------------
NAVY, SLATE, TEAL, ZEBRA, RULE, NOTE = "1F3864", "2F4B7C", "1F7A8C", "EEF2F7", "9AA5B1", "555555"
FONT = "Times New Roman"
OWNER = "MUHAMMED FAIZ"
MONEY = "#,##0.00"

thin = Side(style="thin", color=RULE)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
fill = lambda c: PatternFill("solid", fgColor=c)
NAME_LINE = Font(name=FONT, size=12, bold=True, color=NAVY)
TITLE = Font(name=FONT, size=15, bold=True, color="FFFFFF")
HEADER = Font(name=FONT, size=11, bold=True, color="FFFFFF")
BODY = Font(name=FONT, size=11)
LINK = Font(name=FONT, size=11, color="0563C1", underline="single")
TOTAL = Font(name=FONT, size=12, bold=True, color="FFFFFF")
NOTE_FONT = Font(name=FONT, size=10, italic=True, color=NOTE)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")

# (header, width, alignment)
MONTH_COLS = [("Sl No", 8, CENTER), ("Date", 15, CENTER), ("Name", 40, LEFT),
              ("Service", 50, CENTER), ("Charged", 18, RIGHT)]
SUMMARY_COLS = [("Sl No", 8, CENTER), ("Month", 22, CENTER),
                ("Total Entries", 18, CENTER), ("Total Charged", 22, RIGHT)]


def style_range(ws, row, first, last, font, fill_color=None, align=CENTER, border=True):
    for col in range(first, last + 1):
        c = ws.cell(row=row, column=col)
        c.font, c.alignment = font, align
        if fill_color:
            c.fill = fill(fill_color)
        if border:
            c.border = BORDER


def band(ws, row, text, ncols, font, fill_color, height, align=CENTER):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    ws.cell(row=row, column=1, value=text)
    style_range(ws, row, 1, ncols, font, fill_color, align, border=fill_color is not None)
    ws.row_dimensions[row].height = height


def sheet_frame(ws, title, cols):
    n = len(cols)
    band(ws, 1, f"NAME : {OWNER}", n, NAME_LINE, None, 21.75, LEFT)
    band(ws, 2, title, n, TITLE, NAVY, 27.75)
    for i, (head, width, _) in enumerate(cols, start=1):
        ws.cell(row=3, column=i, value=head)
        ws.column_dimensions[chr(64 + i)].width = width
    style_range(ws, 3, 1, n, HEADER, SLATE)
    ws.row_dimensions[3].height = 30
    ws.freeze_panes = "A4"
    ws.sheet_view.showGridLines = False
    ps = ws.page_setup
    ps.paperSize, ps.orientation = ws.PAPERSIZE_A4, "portrait"
    ps.fitToWidth, ps.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.print_options.horizontalCentered = True
    ws.print_title_rows = "1:3"
    ws.oddFooter.center.text = "Page &P of &N"


def body_row(ws, row, cols, index):
    for col, (_, _, align) in enumerate(cols, start=1):
        c = ws.cell(row=row, column=col)
        c.font, c.alignment, c.border = BODY, align, BORDER
        if index % 2:
            c.fill = fill(ZEBRA)
    ws.row_dimensions[row].height = 18


def total_row(ws, row, label, label_cols, value_col, formula, number_format):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=label_cols)
    ws.cell(row=row, column=1, value=label)
    style_range(ws, row, 1, label_cols, TOTAL, NAVY, RIGHT)
    c = ws.cell(row=row, column=value_col, value=formula)
    c.font, c.fill, c.border, c.number_format = TOTAL, fill(NAVY), BORDER, number_format
    c.alignment = RIGHT if number_format == MONEY else CENTER
    ws.row_dimensions[row].height = 24


def divider(ws, row, ncols):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    for col in range(1, ncols + 1):
        ws.cell(row=row, column=col).fill = fill(TEAL)
    ws.row_dimensions[row].height = 7.5


def month_sheet(wb, year, month, rows):
    label = f"{calendar.month_name[month].upper()} {year}"
    ws = wb.create_sheet(f"{calendar.month_abbr[month]}-{year}")
    sheet_frame(ws, f"TRAVELS DATA OF {label}", MONTH_COLS)
    first, last = 4, 3 + rows
    for i, r in enumerate(range(first, last + 1)):
        body_row(ws, r, MONTH_COLS, i)
        # Sl No numbers itself once a row has a Name.
        ws.cell(row=r, column=1, value=f'=IF(C{r}="","",COUNTA($C${first}:C{r}))')
        ws.cell(row=r, column=2).number_format = "@"      # dates as text: 01 Jan 2026
        ws.cell(row=r, column=5).number_format = MONEY

    dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0",
                        error="Charged must be a number of 0 or more.", errorTitle="Charged")
    dv.add(f"E{first}:E{last}")
    ws.add_data_validation(dv)

    divider(ws, last + 1, 5)
    count_row, sum_row = last + 2, last + 3
    total_row(ws, count_row, f"TOTAL ENTRIES OF {label}", 4, 5,
              f"=COUNTA(C{first}:C{last})", "0")
    total_row(ws, sum_row, f"TOTAL CHARGED OF {label}", 4, 5,
              f"=SUM(E{first}:E{last})", MONEY)
    note = ws.cell(row=sum_row + 2, column=1,
                   value="Fill in Date (as 01 Jan 2026), Name, Service and Charged. "
                         "Sl No and the totals calculate themselves.")
    note.font = NOTE_FONT
    return ws.title, count_row, sum_row


def summary_sheet(ws, year, months):
    ws.title = "Summary"
    sheet_frame(ws, f"TRAVELS SUMMARY: JAN {year} - DEC {year}", SUMMARY_COLS)
    first = 4
    for i, (title, count_row, sum_row) in enumerate(months):
        r = first + i
        body_row(ws, r, SUMMARY_COLS, i)
        ws.cell(row=r, column=1, value=i + 1)
        month = ws.cell(row=r, column=2, value=title)
        month.hyperlink, month.font = f"#'{title}'!A1", LINK
        ws.cell(row=r, column=3, value=f"='{title}'!E{count_row}").number_format = "0"
        ws.cell(row=r, column=4, value=f"='{title}'!E{sum_row}").number_format = MONEY
    last = first + len(months) - 1
    divider(ws, last + 1, 4)
    total_row(ws, last + 2, f"GRAND TOTAL {year}", 2, 3, f"=SUM(C{first}:C{last})", "0")
    c = ws.cell(row=last + 2, column=4, value=f"=SUM(D{first}:D{last})")
    c.font, c.fill, c.border, c.number_format, c.alignment = TOTAL, fill(NAVY), BORDER, MONEY, RIGHT
    ws.cell(row=last + 4, column=1,
            value="Each month's figures are read from the totals at the bottom of its sheet.").font = NOTE_FONT


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--year", type=int, default=2026)
    ap.add_argument("--rows", type=int, default=50, help="blank rows per month sheet")
    ap.add_argument("--out")
    args = ap.parse_args()
    out = args.out or f"output/Faiz_Travels_Jan{args.year}-Dec{args.year}.xlsx"

    wb = Workbook()
    summary = wb.active
    months = [month_sheet(wb, args.year, m, args.rows) for m in range(1, 13)]
    summary_sheet(summary, args.year, months)
    wb.save(out)
    print(out)


if __name__ == "__main__":
    main()
