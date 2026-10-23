"""Yearly workbooks in the Faiz Medical Records house style.

A "Summary" sheet followed by one sheet per month (``Jan-2026`` ...). Each month
sheet lists that month's entries in date order, keeps blank rows below them for
typing by hand, and closes with totals. Sl No, the totals and the Summary are
live formulas; only the entries themselves are written as values.

A section (TRAVELS, PATIENTS, RADIOLOGY) is described by a spec dict, see
``SECTIONS`` at the bottom.
"""
import calendar
import os
import shutil
import subprocess
import tempfile
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# ---- Faiz Medical Records tokens ------------------------------------------
NAVY, SLATE, TEAL, ZEBRA, RULE, NOTE = "1F3864", "2F4B7C", "1F7A8C", "EEF2F7", "9AA5B1", "555555"
FONT = "Times New Roman"
OWNER = "MUHAMMED FAIZ"
MONEY = "#,##0.00"
SPARE_ROWS = 10        # blank rows kept below a busy month's last entry

thin = Side(style="thin", color=RULE)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
fill = lambda c: PatternFill("solid", fgColor=c)
NAME_LINE = Font(name=FONT, size=12, bold=True, color=NAVY)
TITLE = Font(name=FONT, size=15, bold=True, color="FFFFFF")
HEADER = Font(name=FONT, size=11, bold=True, color="FFFFFF")
BODY = Font(name=FONT, size=11)
LINK = Font(name=FONT, size=11, color="0563C1", underline="single")
TOTAL = Font(name=FONT, size=12, bold=True, color="FFFFFF")
NOTE_FONT = Font(name=FONT, size=10, italic=True, color=NOTE)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")


def style_range(ws, row, first, last, font, fill_color=None, align=CENTER, border=True):
    for col in range(first, last + 1):
        c = ws.cell(row=row, column=col)
        c.font, c.alignment = font, align
        if fill_color:
            c.fill = fill(fill_color)
        if border:
            c.border = BORDER


def band(ws, row, text, ncols, font, fill_color, height, align=CENTER):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    ws.cell(row=row, column=1, value=text)
    style_range(ws, row, 1, ncols, font, fill_color, align, border=fill_color is not None)
    ws.row_dimensions[row].height = height


def sheet_frame(ws, title, cols):
    n = len(cols)
    band(ws, 1, f"NAME : {OWNER}", n, NAME_LINE, None, 21.75, LEFT)
    band(ws, 2, title, n, TITLE, NAVY, 27.75)
    for i, (head, width, *_) in enumerate(cols, start=1):
        ws.cell(row=3, column=i, value=head)
        ws.column_dimensions[get_column_letter(i)].width = width
    style_range(ws, 3, 1, n, HEADER, SLATE)
    ws.row_dimensions[3].height = 30
    ws.freeze_panes = "A4"
    ws.sheet_view.showGridLines = False
    ps = ws.page_setup
    ps.paperSize, ps.orientation = ws.PAPERSIZE_A4, "portrait"
    ps.fitToWidth, ps.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.print_options.horizontalCentered = True
    ws.print_title_rows = "1:3"
    ws.oddFooter.center.text = "Page &P of &N"


def body_row(ws, row, cols, index):
    for col, (_, _, align, *_) in enumerate(cols, start=1):
        c = ws.cell(row=row, column=col)
        c.font, c.alignment, c.border = BODY, align, BORDER
        if index % 2:
            c.fill = fill(ZEBRA)
    ws.row_dimensions[row].height = 18


def total_row(ws, row, label, label_cols, value_col, formula, number_format):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=label_cols)
    ws.cell(row=row, column=1, value=label)
    style_range(ws, row, 1, label_cols, TOTAL, NAVY, RIGHT)
    c = ws.cell(row=row, column=value_col, value=formula)
    c.font, c.fill, c.border, c.number_format = TOTAL, fill(NAVY), BORDER, number_format
    c.alignment = RIGHT if number_format == MONEY else CENTER
    ws.row_dimensions[row].height = 24


def divider(ws, row, ncols):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    for col in range(1, ncols + 1):
        ws.cell(row=row, column=col).fill = fill(TEAL)
    ws.row_dimensions[row].height = 7.5


def text_date(iso):
    """2026-10-23 -> 23 Oct 2026, the way the records have always been written."""
    return date.fromisoformat(iso).strftime("%d %b %Y")


def month_sheet(wb, spec, year, month, entries, rows):
    """entries: this month's entries, already in date order. Returns (title, {total key: row})."""
    cols = spec["columns"]
    keys = [c[3] for c in cols]
    ncols = len(cols)
    letter = {k: get_column_letter(i) for i, k in enumerate(keys, start=1) if k}
    count_col = letter[spec["count_key"]]
    label = f"{calendar.month_name[month].upper()} {year}"
    ws = wb.create_sheet(f"{calendar.month_abbr[month]}-{year}")
    sheet_frame(ws, spec["title"].format(month=label), cols)

    first = 4
    last = first + max(rows, len(entries) + SPARE_ROWS) - 1
    for i, r in enumerate(range(first, last + 1)):
        body_row(ws, r, cols, i)
        # Sl No numbers itself once a row has an entry.
        ws.cell(row=r, column=1,
                value=f'=IF({count_col}{r}="","",COUNTA(${count_col}${first}:{count_col}{r}))')
        entry = entries[i] if i < len(entries) else {}
        for col, key in enumerate(keys, start=1):
            if key is None:
                continue
            c = ws.cell(row=r, column=col)
            if key == "date":
                c.number_format = "@"          # dates as text: 01 Jan 2026
                if entry.get("date"):
                    c.value = text_date(entry["date"])
            else:
                c.value = entry.get(key)
                if key in spec["money"]:
                    c.number_format = MONEY

    for key in spec["money"]:
        dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0",
                            error="Enter a number of 0 or more.", errorTitle="Amount")
        dv.add(f"{letter[key]}{first}:{letter[key]}{last}")
        ws.add_data_validation(dv)

    divider(ws, last + 1, ncols)
    total_rows = {}
    for k, (text, key, _) in enumerate(spec["totals"]):
        r = last + 2 + k
        if key == "count":
            formula, fmt = f"=COUNTA({count_col}{first}:{count_col}{last})", "0"
        else:
            formula, fmt = f"=SUM({letter[key]}{first}:{letter[key]}{last})", MONEY
        total_row(ws, r, text.format(month=label), ncols - 1, ncols, formula, fmt)
        total_rows[key] = r
    ws.cell(row=last + 3 + len(spec["totals"]), column=1, value=spec["note"]).font = NOTE_FONT
    return ws.title, total_rows


def summary_sheet(ws, spec, year, months):
    ws.title = "Summary"
    value_col = get_column_letter(len(spec["columns"]))
    cols = [("Sl No", 8, CENTER), ("Month", 22, CENTER)] + [
        (head, 22 if key != "count" else 18, RIGHT if key != "count" else CENTER)
        for _, key, head in spec["totals"]]
    sheet_frame(ws, spec["summary_title"].format(year=year), cols)
    first = 4
    for i, (title, total_rows) in enumerate(months):
        r = first + i
        body_row(ws, r, cols, i)
        ws.cell(row=r, column=1, value=i + 1)
        month = ws.cell(row=r, column=2, value=title)
        month.hyperlink, month.font = f"#'{title}'!A1", LINK
        for j, (_, key, _) in enumerate(spec["totals"]):
            ws.cell(row=r, column=3 + j, value=f"='{title}'!{value_col}{total_rows[key]}").number_format = (
                "0" if key == "count" else MONEY)
    last = first + len(months) - 1
    divider(ws, last + 1, len(cols))
    for j, (_, key, _) in enumerate(spec["totals"]):
        col = get_column_letter(3 + j)
        fmt = "0" if key == "count" else MONEY
        if j == 0:
            total_row(ws, last + 2, f"GRAND TOTAL {year}", 2, 3, f"=SUM({col}{first}:{col}{last})", fmt)
        else:
            c = ws.cell(row=last + 2, column=3 + j, value=f"=SUM({col}{first}:{col}{last})")
            c.font, c.fill, c.border, c.number_format = TOTAL, fill(NAVY), BORDER, fmt
            c.alignment = RIGHT if fmt == MONEY else CENTER
    ws.cell(row=last + 4, column=1,
            value="Each month's figures are read from the totals at the bottom of its sheet.").font = NOTE_FONT


def libreoffice(src, outdir, fmt):
    """Convert src with headless LibreOffice; returns the new file's path, or None without LibreOffice."""
    if not shutil.which("soffice"):
        return None
    with tempfile.TemporaryDirectory() as profile:
        env = dict(os.environ, SAL_USE_VCLPLUGIN="svp")
        subprocess.run(["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--norestore",
                        "--convert-to", fmt, "--outdir", outdir, src],
                       env=env, check=True, capture_output=True, timeout=300)
    return os.path.join(outdir, os.path.splitext(os.path.basename(src))[0] + "." + fmt)


def build(spec, year, entries, out, rows=50):
    """Write the year's workbook for one section. entries: dicts with an ISO 'date'.

    The file is passed through LibreOffice when it is installed, so it carries calculated
    results for viewers that don't calculate (phone and Drive previews).
    """
    by_month = {m: [] for m in range(1, 13)}
    for e in sorted(entries, key=lambda e: e["date"]):      # stable: same-day entries keep their order
        d = date.fromisoformat(e["date"])
        if d.year == year:
            by_month[d.month].append(e)
    wb = Workbook()
    summary = wb.active
    months = [month_sheet(wb, spec, year, m, by_month[m], rows) for m in range(1, 13)]
    summary_sheet(summary, spec, year, months)
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        raw = os.path.join(tmp, os.path.basename(out))
        wb.save(raw)
        calculated = libreoffice(raw, os.path.join(tmp, "lo"), "xlsx")
        shutil.copyfile(calculated or raw, out)
    return out


# (header, width, alignment, entry key) - key None is the Sl No formula column.
SECTIONS = {
    "travels": dict(
        title="TRAVELS DATA OF {month}",
        summary_title="TRAVELS SUMMARY: JAN {year} - DEC {year}",
        columns=[("Sl No", 8, CENTER, None), ("Date", 15, CENTER, "date"), ("Name", 40, LEFT, "name"),
                 ("Service", 50, CENTER, "service"), ("Charged", 18, RIGHT, "charged")],
        count_key="name",
        money=["charged"],
        totals=[("TOTAL ENTRIES OF {month}", "count", "Total Entries"),
                ("TOTAL CHARGED OF {month}", "charged", "Total Charged")],
        note="Fill in Date (as 01 Jan 2026), Name, Service and Charged. "
             "Sl No and the totals calculate themselves.",
        file="Faiz_Travels_Jan{year}-Dec{year}.xlsx",
    ),
    "patients": dict(
        title="PATIENT DATA OF {month}",
        summary_title="PATIENT DATA SUMMARY: JAN {year} - DEC {year}",
        columns=[("Sl No", 7, CENTER, None), ("Date", 14, CENTER, "date"),
                 ("File No.", 18, CENTER, "file_no"), ("Name", 32, LEFT, "name"),
                 ("Consultation", 24, CENTER, "consultation"), ("Services", 24, CENTER, "services"),
                 ("Bill", 14, RIGHT, "bill")],
        count_key="name",
        money=["bill"],
        totals=[("TOTAL PATIENTS OF {month}", "count", "Total Patients"),
                ("TOTAL BILL OF {month}", "bill", "Total Bill")],
        note="Fill in Date (as 01 Jan 2026), File No., Name, Consultation, Services and Bill. "
             "Sl No and the totals calculate themselves.",
        file="Faiz_Patient_Data_Jan{year}-Dec{year}.xlsx",
    ),
    "radiology": dict(
        title="RADIOLOGY DATA OF {month} - OUT PATIENTS",
        summary_title="RADIOLOGY SUMMARY: JAN {year} - DEC {year}",
        columns=[("Sl No", 8, CENTER, None), ("Date of Visit", 15, CENTER, "date"),
                 ("File No.", 20, CENTER, "file_no"), ("Patient", 42, LEFT, "patient"),
                 ("Consultation", 26, CENTER, "consultation")],
        count_key="patient",
        money=[],
        totals=[("TOTAL REFERRALS OF {month}", "count", "Total Referrals")],
        note="Fill in Date of Visit (as 01 Jan 2026), File No., Patient and Consultation. "
             "Sl No and the totals calculate themselves.",
        file="Faiz_Radiology_Data_Jan{year}-Dec{year}.xlsx",
    ),
}
